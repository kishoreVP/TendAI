"""Report model for a run, and JSON / CSV / Excel / PDF exports of it."""
import csv
import io
import json
from collections import Counter

from .db import AuditEvent, Bidder, Document, Evidence, PriceComparison, Result, TemplateVersion


def run_report(s, run):
    tv = s.get(TemplateVersion, run.template_version_id)
    reqs = {r["requirement_id"]: r for r in tv.requirements}
    bidders = {b.id: b.name for b in s.query(Bidder).filter_by(case_id=run.case_id)}
    evidence = {e.id: e.data | {"id": e.id} for e in s.query(Evidence).filter_by(run_id=run.id)}
    results = []
    for r in s.query(Result).filter_by(run_id=run.id).order_by(Result.id):
        results.append({"id": r.id, "bidder_id": r.bidder_id, "bidder": bidders[r.bidder_id],
                        "requirement_key": r.requirement_key, "auto_status": r.auto_status,
                        "final_status": r.final_status, "status": r.final_status or r.auto_status,
                        "reviewed": r.final_status is not None, "reviewer": r.reviewer, "review_reason": r.review_reason,
                        "explanation": r.explanation, "calc": r.calc,
                        "evidence": [evidence[i] for i in r.evidence_ids if i in evidence]})
    prices = [{"bidder_id": p.bidder_id, "bidder": bidders[p.bidder_id], "node_id": p.node_id, **p.data}
              for p in s.query(PriceComparison).filter_by(run_id=run.id)]
    totals = Counter()
    for p in prices:
        totals[p["bidder"]] += p["amount"] or 0
    by_bidder = {name: dict(Counter(r["status"] for r in results if r["bidder"] == name)) for name in bidders.values()}
    docs = s.query(Document).filter_by(case_id=run.case_id).all()
    return {
        "run": {"id": run.id, "created": run.created.isoformat(), "config": run.config},
        "template": {"id": tv.id, "version": tv.version, "approved_by": tv.approved_by,
                     "created": tv.created.isoformat()},
        "requirements": list(reqs.values()),
        "bidders": [{"id": i, "name": n} for i, n in bidders.items()],
        "documents": [{"id": d.id, "filename": d.filename, "role": d.role, "kind": d.kind, "bidder_id": d.bidder_id,
                       "sha256": d.sha256, "warnings": d.warnings} for d in docs],
        "summary": {"requirements": len(reqs), "status_by_bidder": by_bidder,
                    "reviewed_results": sum(r["reviewed"] for r in results),
                    "total_results": len(results),
                    "benchmark_coverage": {
                        "items": len({p["sl"] for p in prices}),
                        "with_comparable_reference": len({p["sl"] for p in prices if p["benchmark"]["comparable"]}),
                        "with_partial_reference": len({p["sl"] for p in prices
                                                       if p["benchmark"]["partial"] and not p["benchmark"]["comparable"]})},
                    "quoted_totals": dict(totals),
                    "extraction_warnings": sum(len(d.warnings) for d in docs),
                    "requirement_flags": sum(bool(r.get("flags")) for r in reqs.values())},
        "results": results,
        "prices": prices,
        "audit": [{"ts": a.ts.isoformat(), "actor": a.actor, "action": a.action, "entity": a.entity,
                   "entity_id": a.entity_id, "detail": a.detail}
                  for a in s.query(AuditEvent).order_by(AuditEvent.id)],
    }


MATRIX_COLS = ["requirement_id", "category", "requirement", "government_value", "bidder", "bidder_value",
               "auto_status", "final_status", "explanation", "evidence", "review"]


def matrix_rows(report):
    reqs = {r["requirement_id"]: r for r in report["requirements"]}
    for r in report["results"]:
        q = reqs[r["requirement_key"]]
        ev = r["evidence"]
        yield {"requirement_id": q["requirement_id"], "category": q["category"],
               "requirement": q["normalised_requirement"],
               "government_value": q.get("expected_value") or q["evaluation_rule"].get("value") or
               ", ".join(q["evaluation_rule"].get("permitted_values") or []),
               "bidder": r["bidder"], "bidder_value": "; ".join(str(e.get("raw_value")) for e in ev),
               "auto_status": r["auto_status"], "final_status": r["final_status"] or "",
               "explanation": r["explanation"],
               "evidence": "; ".join(f"doc {e['source_document_id']} p.{e.get('source_page')} {e.get('source_label') or ''}"
                                     for e in ev),
               "review": f"{r['reviewer']}: {r['review_reason']}" if r["reviewed"] else "pending human review"}


def to_json(s, run):
    return json.dumps(run_report(s, run), indent=2, default=str).encode(), "application/json"


def to_csv(s, run):
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=MATRIX_COLS)
    w.writeheader()
    w.writerows(matrix_rows(run_report(s, run)))
    return buf.getvalue().encode("utf-8-sig"), "text/csv"


def to_xlsx(s, run):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    rep = run_report(s, run)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Compliance matrix"
    ws.append(MATRIX_COLS)
    fills = {"MEETS_REQUIREMENT": "C6EFCE", "DOES_NOT_MEET_REQUIREMENT": "FFC7CE", "INSUFFICIENT_EVIDENCE": "FFEB9C",
             "REQUIRES_HUMAN_REVIEW": "DDEBF7"}
    for row in matrix_rows(rep):
        ws.append([row[c] for c in MATRIX_COLS])
        c = ws.cell(ws.max_row, MATRIX_COLS.index("auto_status") + 1)
        if c.value in fills:
            c.fill = PatternFill("solid", fgColor=fills[c.value])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws2 = wb.create_sheet("Commercial")
    cols = ["bidder", "sl", "description", "qty", "unit", "quoted_rate", "amount", "benchmark_median",
            "ratio_to_benchmark", "comparable_refs", "partial_refs"]
    ws2.append(cols)
    for p in rep["prices"]:
        b = p["benchmark"]
        ws2.append([p["bidder"], p["sl"], p["description"][:200], p["qty"], p["unit"], p["quoted_rate"], p["amount"],
                    (b.get("summary") or {}).get("median"), p.get("ratio_to_benchmark"),
                    "; ".join(f"{e['code']} {e['normalised_rate']}" for e in b["comparable"][:3]),
                    "; ".join(f"{e['code']} {e['normalised_rate']} ({'; '.join(e['notes'])})" for e in b["partial"][:3])])
    ws3 = wb.create_sheet("Template")
    ws3.append(["requirement_id", "category", "title", "operator", "value", "unit", "mandatory", "source_page",
                "source_clause", "flags", "original_clause_text"])
    for q in rep["requirements"]:
        ws3.append([q["requirement_id"], q["category"], q["title"], q["comparison_operator"],
                    str(q["evaluation_rule"].get("value") or q["evaluation_rule"].get("permitted_values") or ""),
                    q.get("unit"), q["mandatory_status"], q.get("source_page"), q.get("source_clause"),
                    "; ".join(q.get("flags", [])), q["original_clause_text"][:1000]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def to_pdf(s, run):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    rep = run_report(s, run)
    st = getSampleStyleSheet()
    small = st["BodyText"].clone("small", fontSize=7, leading=9)
    esc = lambda t: str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    story = [Paragraph(f"TendAI evaluation report - run {run.id}", st["Title"]),
             Paragraph("Decision-support output. Statuses marked 'pending human review' are proposals, "
                       "not authorised decisions.", st["Italic"]),
             Paragraph(f"Template version {rep['template']['version']} approved by {esc(rep['template']['approved_by'])}; "
                       f"rule engine {rep['run']['config']['rule_engine_version']}; created {rep['run']['created']}",
                       st["BodyText"]), Spacer(1, 8),
             Paragraph("Source documents", st["Heading2"])]
    story.append(Table([["id", "file", "role", "sha256"]] +
                       [[d["id"], Paragraph(esc(d["filename"]), small), d["role"], d["sha256"][:16]]
                        for d in rep["documents"]], colWidths=[30, 330, 60, 120]))
    story.append(Paragraph("Status by bidder", st["Heading2"]))
    for b, counts in rep["summary"]["status_by_bidder"].items():
        story.append(Paragraph(f"<b>{esc(b)}</b>: " + ", ".join(f"{k}: {v}" for k, v in counts.items()), st["BodyText"]))
    story.append(Paragraph("Compliance matrix", st["Heading2"]))
    rows = [["Req", "Bidder", "Bidder value", "Status", "Explanation", "Evidence"]]
    for m in matrix_rows(rep):
        rows.append([Paragraph(esc(m["requirement_id"]), small), Paragraph(esc(m["bidder"]), small),
                     Paragraph(esc(m["bidder_value"])[:120], small), Paragraph(esc(m["final_status"] or m["auto_status"]), small),
                     Paragraph(esc(m["explanation"])[:400], small), Paragraph(esc(m["evidence"])[:150], small)])
    t = Table(rows, colWidths=[60, 70, 90, 90, 330, 130], repeatRows=1)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey)]))
    story.append(t)
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=20, rightMargin=20).build(story)
    return buf.getvalue(), "application/pdf"


# ------------------------------------------------------- visual evidence
def snippet_png(path, page_no, bbox, pad=14, dpi=130):
    """PNG of a PDF region with the region itself highlighted - the visual proof of a finding.
    bbox is in displayed-page coordinates (as stored on nodes)."""
    import pymupdf
    doc = pymupdf.open(path)
    page = doc[page_no - 1]
    rect = pymupdf.Rect(bbox) * ~page.rotation_matrix  # back to unrotated page space
    page.draw_rect(rect, color=(0.95, 0.65, 0), fill=(1, 0.85, 0.2), fill_opacity=0.28, width=1.2)
    clip = pymupdf.Rect(rect.x0 - pad, rect.y0 - pad, rect.x1 + pad, rect.y1 + pad) & page.cropbox
    if page.rotation == 0:  # full width keeps the table / sentence context around the value
        clip.x0, clip.x1 = page.cropbox.x0 + 10, page.cropbox.x1 - 10
    png = page.get_pixmap(clip=clip, dpi=dpi).tobytes("png")
    doc.close()
    return png


def first_pdf_region(regions):
    return next((g for g in regions or [] if g.get("page") and g.get("bbox")), None)


def evidence_pack(s, run, bidder_id):
    """Per-bidder PDF: summary, then every finding that is not a clean pass, each with the
    tender clause and the bidder's evidence rendered from the source documents."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table

    rep = run_report(s, run)
    bidder = next(b["name"] for b in rep["bidders"] if b["id"] == bidder_id)
    reqs = {r["requirement_id"]: r for r in rep["requirements"]}
    docs = {d["id"]: d for d in rep["documents"]}
    st = getSampleStyleSheet()
    esc = lambda t: str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    colour = {"MEETS_REQUIREMENT": "#17602f", "DOES_NOT_MEET_REQUIREMENT": "#a02622",
              "INSUFFICIENT_EVIDENCE": "#7a5600", "REQUIRES_HUMAN_REVIEW": "#204a97"}

    def image(doc_id, regions, width):
        g = first_pdf_region(regions)
        doc = s.get(Document, doc_id)
        if not g or doc is None or doc.kind != "pdf":
            return Paragraph("<i>no PDF region (spreadsheet / not located)</i>", st["BodyText"])
        img = Image(io.BytesIO(snippet_png(doc.path, g["page"], g["bbox"])))
        img.drawWidth, img.drawHeight = width, width * img.imageHeight / img.imageWidth
        return img

    mine = [r for r in rep["results"] if r["bidder"] == bidder]
    counts = Counter(r["status"] for r in mine)
    story = [Paragraph(f"TendAI evidence pack - {esc(bidder)}", st["Title"]),
             Paragraph(f"Run {run.id}, template v{rep['template']['version']}. Proposed statuses; final only after "
                       "a named reviewer records a decision.", st["Italic"]), Spacer(1, 6),
             Paragraph(" &nbsp; ".join(f"<font color='{colour.get(k, '#333')}'><b>{v}</b> {k.replace('_', ' ').lower()}</font>"
                                       for k, v in counts.items()), st["BodyText"]), Spacer(1, 10)]
    order = ["DOES_NOT_MEET_REQUIREMENT", "REQUIRES_HUMAN_REVIEW", "INSUFFICIENT_EVIDENCE"]
    for r in sorted([r for r in mine if r["status"] in order], key=lambda r: order.index(r["status"])):
        q = reqs[r["requirement_key"]]
        ev = r["evidence"][0] if r["evidence"] else None
        head = Paragraph(f"<font color='{colour[r['status']]}'><b>{r['status'].replace('_', ' ')}</b></font> - "
                         f"<b>{esc(q['requirement_id'])}</b> {esc(q['title'])}", st["Heading4"])
        why = Paragraph(esc(r["explanation"])[:600], st["BodyText"])
        tender = [Paragraph(f"<b>Tender</b> - {esc(docs[q['source_document_id']]['filename'])} p.{q.get('source_page')}",
                            st["BodyText"]), image(q["source_document_id"], q["source_regions"], 250)]
        bid = [Paragraph(f"<b>Bidder</b> - {esc(docs[ev['source_document_id']]['filename'])} p.{ev.get('source_page')}"
                         if ev else "<b>Bidder</b> - no evidence found", st["BodyText"]),
               image(ev["source_document_id"], ev["source_regions"], 250) if ev else Spacer(1, 1)]
        grid = Table([[tender, bid]], colWidths=[262, 262])
        grid.setStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOX", (0, 0), (-1, -1), 0.4, colors.lightgrey)])
        story.append(KeepTogether([head, why, Spacer(1, 4), grid, Spacer(1, 12)]))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=30, rightMargin=30, title=f"Evidence pack {bidder}").build(story)
    return buf.getvalue()
