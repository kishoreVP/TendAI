"""Milestones 5, 6 and 8: full workflow on the sample tenders + synthetic bidders, checked
against the manually written golden outcomes (tests/golden/expected_outcomes.json)."""
import io
import json

import openpyxl
import pymupdf
import pytest

from tendai import config, exports, pipeline
from tendai.db import Document, Node, Requirement, Result, make_session_factory

from .conftest import CABLE, GOLDEN, HIST, PWD, needs_samples

pytestmark = needs_samples


@pytest.fixture(scope="module")
def world(tmp_path_factory, sor_extract):
    old = config.STORAGE_DIR
    config.STORAGE_DIR = tmp_path_factory.mktemp("storage")
    s = make_session_factory("sqlite://")()
    runs = {}
    meta = {"issuer": "TN PWD", "effective_date": "2025-07-22", "tax_basis": "excludes GST (SOR note xi)"}
    for name, folder, bench in [("cable_m246", CABLE, [HIST]), ("pwd_tiruchengode", PWD, [sor_extract])]:
        case = pipeline.load_sample_case(s, folder, name, bench, meta)
        pipeline.extract_requirements(s, case, use_llm=False)
        pipeline.extract_benchmarks(s, case)
        for r in s.query(Requirement).filter_by(case_id=case.id):
            pipeline.review_requirement(s, r, "test", "approved")
        tv = pipeline.approve_template(s, case, "test")
        runs[name] = (case, tv, pipeline.run_evaluation(s, case, tv, use_llm=False))
    yield s, runs
    s.close()
    config.STORAGE_DIR = old


def _results(s, run):
    rep = exports.run_report(s, run)
    titles = {r["requirement_id"]: r["title"].lower() for r in rep["requirements"]}
    return rep, [(r, titles[r["requirement_key"]]) for r in rep["results"]]


@pytest.mark.parametrize("case", ["cable_m246", "pwd_tiruchengode"])
def test_matches_golden_outcomes(world, case):
    s, runs = world
    _, results = _results(s, runs[case][2])
    checked, wrong = 0, []
    for bidder, expect in GOLDEN[case].items():
        for key, status in expect.items():
            hits = [r for r, title in results if r["bidder"] == bidder and key in title]
            assert hits, f"no requirement title contains '{key}'"
            for r in hits:
                checked += 1
                if r["auto_status"] != status:
                    wrong.append(f"{bidder}/{r['requirement_key']}: got {r['auto_status']}, expected {status}")
    assert not wrong, "\n".join(wrong)
    assert checked >= 8


def test_evidence_highlights_contain_the_value(world):
    """Visual grounding: the highlighted PDF region of each numeric claim contains its raw value."""
    s, runs = world
    rep, results = _results(s, runs["cable_m246"][2])
    checked = 0
    for r, _ in results:
        for e in r["evidence"]:
            if not isinstance(e["value"], float) or not e["source_regions"] or not e["source_regions"][0].get("bbox"):
                continue
            doc = s.get(Document, e["source_document_id"])
            g = e["source_regions"][0]
            page = pymupdf.open(doc.path)[g["page"] - 1]
            text = page.get_textbox(pymupdf.Rect(g["bbox"]) * ~page.rotation_matrix)
            assert e["raw_value"].split()[0].replace("Rs.", "").strip() in text.replace("\n", " "), (e["raw_value"], text)
            checked += 1
    assert checked >= 20


def test_requirement_sources_point_at_real_pages(world):
    s, runs = world
    for case, tv, run in runs.values():
        for q in tv.requirements:
            doc = s.get(Document, q["source_document_id"])
            n_pages = len(doc.meta.get("pages", [])) or None
            for g in q["source_regions"]:
                if g.get("page"):
                    assert 1 <= g["page"] <= n_pages
                    w, h = doc.meta["pages"][g["page"] - 1]["w"], doc.meta["pages"][g["page"] - 1]["h"]
                    x0, y0, x1, y1 = g["bbox"]
                    assert 0 <= x0 < x1 <= w + 1 and 0 <= y0 < y1 <= h + 1


def test_rerun_is_reproducible_and_keeps_human_decisions(world):
    s, runs = world
    case, tv, run1 = runs["cable_m246"]
    first = s.query(Result).filter_by(run_id=run1.id).first()
    first.final_status, first.reviewer, first.review_reason = "NOT_APPLICABLE", "alice", "test decision"
    run2 = pipeline.run_evaluation(s, case, tv, use_llm=False)
    a = {(r.bidder_id, r.requirement_key): r.auto_status for r in s.query(Result).filter_by(run_id=run1.id)}
    b = {(r.bidder_id, r.requirement_key): r.auto_status for r in s.query(Result).filter_by(run_id=run2.id)}
    assert a == b
    assert s.get(Result, first.id).final_status == "NOT_APPLICABLE"  # old run untouched
    assert all(r.final_status is None for r in s.query(Result).filter_by(run_id=run2.id))
    assert run2.config["documents"] == run1.config["documents"]


def test_prompt_injection_text_has_no_effect(world):
    s, runs = world
    _, results = _results(s, runs["cable_m246"][2])
    vajra = [r for r, _ in results if r["bidder"] == "vajra"]
    assert any(r["auto_status"] == "DOES_NOT_MEET_REQUIREMENT" for r in vajra)
    injected = s.query(Node).filter(Node.text.contains("ignore all previous instructions")).all()
    assert injected  # the text was ingested as data ...
    assert all(r["auto_status"] != "MEETS_REQUIREMENT" for r in vajra if r["requirement_key"] in ("EMD", "EXPERIENCE-QTY"))


def test_commercial_comparison(world):
    s, runs = world
    rep, _ = _results(s, runs["cable_m246"][2])
    cable = [p for p in rep["prices"] if p["bidder"] == "coromandel"][0]
    bm = cable["benchmark"]
    assert {e["code"] for e in bm["comparable"]} == {"H-001", "H-002"}  # 630 sq.mm, 33 kV only
    assert {e["code"] for e in bm["excluded"]} == {"H-003", "H-004"}
    per_m = next(e for e in bm["comparable"] if e["code"] == "H-002")
    assert per_m["normalised_rate"] == 2150000 and "unit conversion" in per_m["steps"][0]
    assert cable["quoted_rate"] == 2140000 and cable["ratio_to_benchmark"] == round(2140000 / 2120000, 3)
    rep2, _ = _results(s, runs["pwd_tiruchengode"][2])
    mcb = next(p for p in rep2["prices"] if p["sl"] == "16" and p["bidder"] == "murugan")["benchmark"]
    codes = {e["code"] for e in mcb["partial"]}
    assert {"M-1765", "M-1777"} <= codes  # 2-pole 6-32 A MCB, higher and lower end
    assert not mcb["comparable"]  # supply-and-fix item vs material-only rate -> only partial
    assert all(e["code"] not in ("M-1766", "M-1771") for e in mcb["partial"])  # 40 A / 4-pole excluded


def test_exports(world):
    s, runs = world
    run = runs["pwd_tiruchengode"][2]
    data, _ = exports.to_xlsx(s, run)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Compliance matrix", "Commercial", "Template"]
    n_bidders = len(exports.run_report(s, run)["bidders"])
    assert wb["Compliance matrix"].max_row == 1 + len(runs["pwd_tiruchengode"][1].requirements) * n_bidders
    pdf, _ = exports.to_pdf(s, run)
    assert pdf.startswith(b"%PDF") and pymupdf.open(stream=pdf).page_count >= 1
    csv_bytes, _ = exports.to_csv(s, run)
    assert b"pending human review" in csv_bytes
    rep = json.loads(exports.to_json(s, run)[0])
    assert rep["run"]["config"]["rule_engine_version"] and rep["audit"]


def test_snapshots_and_evidence_pack(world):
    """Visual grounding outside the viewer: highlighted crops and the per-bidder evidence pack."""
    s, runs = world
    rep, _ = _results(s, runs["cable_m246"][2])
    meridian = next(b["id"] for b in rep["bidders"] if b["name"] == "meridian")
    r = next(r for r in rep["results"] if r["bidder"] == "meridian" and r["requirement_key"] == "GTP-6")
    e = r["evidence"][0]
    g = exports.first_pdf_region(e["source_regions"])
    png = exports.snippet_png(s.get(Document, e["source_document_id"]).path, g["page"], g["bbox"])
    assert png.startswith(b"\x89PNG")
    pack = pymupdf.open(stream=exports.evidence_pack(s, runs["cable_m246"][2], meridian))
    assert "GTP-6" in "".join(p.get_text() for p in pack)
    assert sum(len(p.get_images()) for p in pack) >= 8  # tender + bidder crop per finding


def test_document_evidence_covers_whole_section(world):
    s, runs = world
    rep, _ = _results(s, runs["pwd_tiruchengode"][2])
    r = next(r for r in rep["results"] if r["bidder"] == "murugan" and r["requirement_key"] == "DOC-INDEMNITY-BOND")
    box = r["evidence"][0]["source_regions"][0]["bbox"]
    assert box[3] - box[1] > 40  # heading + body text, not just the heading line
