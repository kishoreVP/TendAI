"""File validation, content-addressed storage and conversion into Node rows."""
import csv
import hashlib
import re
import shutil
import time
from pathlib import Path

from . import config, pdfparse
from .db import Bidder, Case, Document, Node, audit

MAGIC = {  # extension -> accepted leading bytes
    "pdf": [b"%PDF"],
    "xlsx": [b"PK\x03\x04"], "docx": [b"PK\x03\x04"],
    "xls": [b"\xd0\xcf\x11\xe0"],
    "csv": None,  # plain text; checked by decoding
}


class IngestError(ValueError):
    pass


def safe_name(name):
    """Strip any directory part and unusual characters (no path traversal)."""
    name = Path(name.replace("\\", "/")).name
    return re.sub(r"[^A-Za-z0-9._ -]", "_", name)[:200] or "upload"


def validate(path, filename):
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in MAGIC:
        raise IngestError(f"Unsupported file type: .{ext}")
    size = Path(path).stat().st_size
    if size == 0:
        raise IngestError("File is empty")
    if size > config.MAX_UPLOAD_MB * 1024 * 1024:
        raise IngestError(f"File larger than {config.MAX_UPLOAD_MB} MB")
    head = Path(path).read_bytes()[:8]
    if MAGIC[ext] and not any(head.startswith(m) for m in MAGIC[ext]):
        raise IngestError(f"File content does not look like .{ext}")
    if ext == "csv":
        try:
            Path(path).read_text(encoding="utf-8-sig")
        except UnicodeDecodeError:
            raise IngestError("CSV is not valid UTF-8 text")
    return ext


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def get_case(s, name):
    case = s.query(Case).filter_by(name=name).one_or_none()
    if case is None:
        case = Case(name=name)
        s.add(case)
        s.flush()
    return case


def get_bidder(s, case, name):
    b = s.query(Bidder).filter_by(case_id=case.id, name=name).one_or_none()
    if b is None:
        b = Bidder(case_id=case.id, name=name)
        s.add(b)
        s.flush()
    return b


def ingest(s, case, src_path, role, bidder=None, filename=None, actor="system"):
    """Validate, store and parse one file. Re-ingesting identical bytes returns the existing
    document (idempotent retries)."""
    filename = safe_name(filename or Path(src_path).name)
    kind = validate(src_path, filename)
    digest = sha256(src_path)
    existing = s.query(Document).filter_by(case_id=case.id, sha256=digest, role=role,
                                           bidder_id=bidder.id if bidder else None).one_or_none()
    if existing and existing.status == "parsed":
        return existing
    dest = config.STORAGE_DIR / "files" / f"{digest}.{kind}"
    dest.parent.mkdir(parents=True, exist_ok=True)
    if not dest.exists():
        shutil.copyfile(src_path, dest)
    doc = existing or Document(case_id=case.id, bidder_id=bidder.id if bidder else None, role=role,
                               filename=filename, sha256=digest, path=str(dest), kind=kind)
    s.add(doc)
    s.flush()
    s.query(Node).filter_by(doc_id=doc.id).delete()  # clean slate when retrying a failed parse
    t0 = time.perf_counter()
    try:
        if kind == "pdf":
            nodes, meta = _parse_pdf(dest)
        elif kind in ("xls", "xlsx", "csv"):
            nodes, meta = _parse_sheet(dest, kind)
        else:
            nodes, meta = _parse_docx(dest)
    except IngestError:
        raise
    except Exception as e:  # corrupted file etc.: keep the record, mark it failed
        doc.status, doc.warnings = "failed", [f"parse error: {e}"]
        audit(s, actor, "ingest_failed", "document", doc.id, filename=filename, error=str(e))
        s.flush()
        raise IngestError(f"Could not parse {filename}: {e}") from e
    _save_nodes(s, doc, nodes)
    doc.meta = {**meta, "parse_seconds": round(time.perf_counter() - t0, 2), "parser": "pymupdf+gammaprime"}
    doc.warnings = [f"page {n['page']}: {n['text']}" for n in nodes if n["kind"] == "warning"]
    doc.status = "parsed"
    audit(s, actor, "ingest", "document", doc.id, filename=filename, sha256=digest, role=role, nodes=len(nodes))
    s.flush()
    return doc


def _save_nodes(s, doc, nodes):
    """Bulk insert with pre-assigned ids so parent links resolve without a flush per node."""
    from sqlalchemy import func
    base = (s.query(func.max(Node.id)).scalar() or 0) + 1
    s.bulk_insert_mappings(Node, [
        dict(id=base + i, doc_id=doc.id, parent_id=base + n["parent"] if n.get("parent") is not None else None,
             kind=n["kind"], ref=n.get("ref"), text=n["text"], page=n.get("page"), regions=n["regions"],
             meta=n.get("meta", {}), seq=i)
        for i, n in enumerate(nodes)])


def _parse_pdf(path):
    pages, tables = pdfparse.extract_pages(path)
    nodes = pdfparse.build_tree(pages, tables)
    return nodes, {"pages": [{"w": p["w"], "h": p["h"]} for p in pages]}


# ------------------------------------------------------------------ sheets
def _col_letter(c):  # 1-based
    s = ""
    while c:
        c, r = divmod(c - 1, 26)
        s = chr(65 + r) + s
    return s


def _read_grid(path, kind):
    """-> {sheet_name: [[value, ...], ...]} with displayed values."""
    if kind == "xls":
        import xlrd
        book = xlrd.open_workbook(str(path))
        return {sh.name: [sh.row_values(r) for r in range(sh.nrows)] for sh in book.sheets()}
    if kind == "xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
    with open(path, encoding="utf-8-sig", newline="") as f:
        return {"csv": list(csv.reader(f))}


def _fmt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _parse_sheet(path, kind):
    """Each non-empty row becomes a node with its cells keyed by column letter. BOQ-style
    continuation rows ("49.01" following item "49" with no quantity) become children of
    their item row so long descriptions stay together."""
    nodes, meta = [], {"sheets": []}
    for sheet, grid in _read_grid(path, kind).items():
        grid = [[_fmt(v) for v in row[:60]] for row in grid]  # NIC BOQs carry macro data far right
        sheet_idx = len(nodes)
        nodes.append({"kind": "sheet", "text": sheet, "ref": sheet, "page": None, "parent": None,
                      "regions": [{"sheet": sheet}], "meta": {}})
        header_row, header = None, {}
        last_item = None
        for r, row in enumerate(grid, start=1):
            cells = {_col_letter(c): v for c, v in enumerate(row, start=1) if v}
            if not cells:
                continue
            joined = " ".join(cells.values()).lower()
            if header_row is None and "description" in joined and ("quantity" in joined or "unit" in joined or "rate" in joined):
                header_row, header = r, cells
                nodes[sheet_idx]["meta"] = {"header_row": r, "header": header}
            first = row[0]
            sub = re.fullmatch(r"(\d+)\.\d+", first)
            is_continuation = (last_item is not None and sub and not cells.get("D")
                               and sub.group(1) == nodes[last_item]["meta"]["sl"])
            idx = len(nodes)
            nodes.append({"kind": "row", "text": " | ".join(cells.values()), "ref": f"{sheet}!R{r}",
                          "page": None, "parent": last_item if is_continuation else sheet_idx,
                          "regions": [{"sheet": sheet, "row": r}],
                          "meta": {"cells": cells, "row": r, "sl": first, "is_header": r == header_row}})
            if header_row and r > header_row and re.fullmatch(r"\d+", first):
                last_item = idx
        meta["sheets"].append({"name": sheet, "rows": len(grid)})
    return nodes, meta


# ------------------------------------------------------------------- docx
def _parse_docx(path):
    """Paragraphs are run through the same clause segmentation as PDF rows; DOCX has no
    reliable page positions, so regions point at paragraph / table indices instead."""
    import docx
    d = docx.Document(str(path))
    pages = [{"page": 1, "w": 0, "h": 0, "has_text": True, "rows": []}]
    for i, p in enumerate(d.paragraphs):
        if p.text.strip():
            bold = bool(p.runs) and all(r.bold for r in p.runs if r.text.strip())
            pages[0]["rows"].append({"text": p.text.strip(), "bbox": [0, i, 0, i], "size": 10, "bold": bold,
                                     "merged_from": 1, "furniture": False, "paragraph": i})
    nodes = pdfparse.build_tree(pages, [])
    for n in nodes:
        n["page"] = None
        n["regions"] = [{"paragraph": int(g["bbox"][1])} for g in n["regions"] if g.get("bbox")]
    for t_i, t in enumerate(d.tables):
        rows = [[c.text.strip() for c in r.cells] for r in t.rows]
        if not rows:
            continue
        tnode = len(nodes)
        nodes.append({"kind": "table", "text": " | ".join(rows[0]), "page": None, "parent": None,
                      "regions": [{"table": t_i}], "ref": f"table {t_i + 1}", "meta": {"header": rows[0]}})
        for r_i, cells in enumerate(rows[1:], start=1):
            nodes.append({"kind": "row", "text": " | ".join(c for c in cells if c), "page": None, "parent": tnode,
                          "regions": [{"table": t_i, "row": r_i}], "ref": cells[0],
                          "meta": {"cells": dict(zip(rows[0], cells))}})
    return nodes, {"paragraphs": len(d.paragraphs), "tables": len(d.tables)}
